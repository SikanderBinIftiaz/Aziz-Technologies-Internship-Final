"""Extract text from uploaded documents. Every parser returns a list of pages (or sections)."""
import logging
import re
from dataclasses import dataclass
from pathlib import Path

from app.core.errors import AppError

log = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
SECTION_CHARS = 3000  # DOCX/TXT/MD have no real pages, so they are split into sections of about this size


@dataclass
class ParsedPage:
    page: int
    text: str


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _paginate(paragraphs: list[str], limit: int = SECTION_CHARS) -> list[ParsedPage]:
    pages: list[ParsedPage] = []
    buffer: list[str] = []
    size = 0
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if buffer and size + len(para) > limit:
            pages.append(ParsedPage(len(pages) + 1, "\n\n".join(buffer)))
            buffer, size = [], 0
        buffer.append(para)
        size += len(para) + 2
    if buffer:
        pages.append(ParsedPage(len(pages) + 1, "\n\n".join(buffer)))
    return pages


def _parse_pdf(path: Path) -> list[ParsedPage]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise AppError("PDF_ENCRYPTED", "This PDF is password-protected.", 422) from exc
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:
            pages.append(ParsedPage(number, text))
    return pages


def _parse_docx(path: Path) -> list[ParsedPage]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    paragraphs = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return _paginate([_clean(p) for p in paragraphs])


def _strip_markdown(text: str) -> str:
    """Turn Markdown headings into plain sentences so '## Annual Leave' does not leak into passages."""
    text = re.sub(r"^#{1,6}[ \t]+(.+?)[ \t#]*$", lambda m: m.group(1) if m.group(1)[-1] in ".!?:" else m.group(1) + ".", text, flags=re.M)
    text = re.sub(r"(\*\*|__)(.+?)\1", r"\2", text)
    return re.sub(r"`([^`\n]+)`", r"\1", text)


def _parse_text(path: Path, markdown: bool = False) -> list[ParsedPage]:
    text = _clean(path.read_text(encoding="utf-8", errors="replace"))
    if markdown:
        text = _strip_markdown(text)
    return _paginate(re.split(r"\n\s*\n", text))


def parse_file(path: Path) -> list[ParsedPage]:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            pages = _parse_pdf(path)
        elif ext == ".docx":
            pages = _parse_docx(path)
        elif ext in {".txt", ".md"}:
            pages = _parse_text(path, markdown=(ext == ".md"))
        else:
            raise AppError("UNSUPPORTED_TYPE", f"Unsupported file type '{ext}'.", 415)
    except AppError:
        raise
    except Exception as exc:
        log.warning("Failed to parse %s: %s", path.name, exc)
        raise AppError(
            "PARSE_FAILED", "The file could not be read. It may be corrupted or not a valid document.", 422
        ) from exc

    if not pages:
        raise AppError(
            "NO_TEXT",
            "No readable text was found. Scanned documents need OCR, which is not supported yet.",
            422,
        )
    return pages
